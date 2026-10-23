"""Builds docs/qa/blackbox-test-log.xlsx from docs/qa/blackbox-rounds.json (T-105).

The JSON is the record of every black-box round: the tester's scores, every issue they
reported, and what was done about each. This renders it for reading:

- Summary: scores per area per round (below 8 tinted ember, 8+ tinted moss) and each
  round's lowest score
- All issues: every issue across rounds, filterable, with its fix and status
- Round N: one sheet per round with its scores, issues and fixes

    python docs/qa/build_blackbox_log.py      (needs openpyxl)
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
SOURCE = HERE / "blackbox-rounds.json"
TARGET = HERE / "blackbox-test-log.xlsx"

# Bunko's palette (DESIGN.md): paper neutrals, moss for good, ember for attention.
HEADER = PatternFill("solid", fgColor="E8DFCF")
GOOD = PatternFill("solid", fgColor="D6E4D8")
LOW = PatternFill("solid", fgColor="F3D9C7")
THIN = Side(style="thin", color="D8CBAF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
BOLD = Font(bold=True, color="26221D")

AREAS = [
    "first_impression_onboarding",
    "library_adding_finding_books",
    "book_page_and_details",
    "reading_experience",
    "tracking_journal_statistics",
    "shelves_organizing",
    "notes_highlights_reviews",
    "assistant_chat",
    "look_and_feel",
    "speed_reliability",
    "overall",
]
ISSUE_COLUMNS = [
    ("Round", 8), ("ID", 9), ("Area", 24), ("Severity", 10), ("Title", 36), ("Steps", 46),
    ("Expected", 36), ("Actual", 46), ("Root cause", 40), ("Fix", 52), ("Files", 40),
    ("Commit", 10), ("Status", 12), ("Verified in round", 10),
]


def label(area: str) -> str:
    return area.replace("_", " ").capitalize()


def header_row(ws, values, widths=None) -> None:
    ws.append(values)
    for i, cell in enumerate(ws[ws.max_row], start=1):
        cell.fill, cell.font, cell.border, cell.alignment = HEADER, BOLD, BORDER, WRAP
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = ws.cell(row=ws.max_row + 1, column=1)


def issue_values(round_no: int, issue: dict) -> list:
    return [
        round_no, issue.get("id"), label(issue.get("area", "")), issue.get("severity"), issue.get("title"),
        issue.get("steps"), issue.get("expected"), issue.get("actual"), issue.get("root_cause", ""),
        issue.get("fix", ""), ", ".join(issue.get("files", [])), issue.get("commit", ""),
        issue.get("status", "Open"), issue.get("verified_in", ""),
    ]


def shown(score):
    """An area the tester couldn't reach that round is 'not tested', never a blank or a zero."""
    return "not tested" if score is None else score


def build() -> None:
    data = json.loads(SOURCE.read_text(encoding="utf-8"))
    rounds = data["rounds"]
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    summary.append(["Bunko black-box usability rounds (T-105). Target: every score at least 8/10."])
    summary["A1"].font = Font(bold=True, size=13, color="26221D")
    summary.append([])
    header_row(summary, ["Area", *[f"Round {r['round']}" for r in rounds]], [34, *[11] * len(rounds)])
    for area in AREAS:
        summary.append([label(area), *[shown(r["scores"].get(area, {}).get("score")) for r in rounds]])
        for cell in summary[summary.max_row][1:]:
            cell.border, cell.alignment = BORDER, Alignment(horizontal="center")
            if isinstance(cell.value, (int, float)):
                cell.fill = GOOD if cell.value >= 8 else LOW
        summary[summary.max_row][0].border = BORDER
    summary.append(["Lowest score", *[min(s["score"] for s in r["scores"].values() if s["score"] is not None) for r in rounds]])
    summary.append(["Issues reported", *[len(r["issues"]) for r in rounds]])
    summary.append(["Issues fixed", *[sum(1 for i in r["issues"] if i.get("status") == "Fixed") for r in rounds]])
    for row in summary.iter_rows(min_row=summary.max_row - 2, max_row=summary.max_row):
        row[0].font = BOLD
        for cell in row:
            cell.border = BORDER
    summary.append([])
    for r in rounds:
        summary.append([f"Round {r['round']} ({r.get('date', '')})", r.get("environment", "")])
        summary[summary.max_row][0].font = BOLD
        summary[summary.max_row][1].alignment = WRAP
    summary.column_dimensions["B"].width = max(summary.column_dimensions["B"].width or 11, 11)

    all_issues = wb.create_sheet("All issues")
    header_row(all_issues, [c for c, _ in ISSUE_COLUMNS], [w for _, w in ISSUE_COLUMNS])
    for r in rounds:
        for issue in r["issues"]:
            all_issues.append(issue_values(r["round"], issue))
            for cell in all_issues[all_issues.max_row]:
                cell.alignment, cell.border = WRAP, BORDER
    all_issues.auto_filter.ref = all_issues.dimensions

    for r in rounds:
        ws = wb.create_sheet(f"Round {r['round']}")
        ws.append([f"Round {r['round']}", r.get("date", ""), r.get("tester", "")])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([r.get("environment", "")])
        ws.append([])
        header_row(ws, ["Area", "Score", "Tester's reason"], [30, 8, 90])
        for area in AREAS:
            s = r["scores"].get(area, {})
            ws.append([label(area), shown(s.get("score")), s.get("reason", "")])
            row = ws[ws.max_row]
            for cell in row:
                cell.border, cell.alignment = BORDER, WRAP
            if isinstance(row[1].value, (int, float)):
                row[1].fill = GOOD if row[1].value >= 8 else LOW
        ws.append([])
        ws.append(["Issues and fixes"])
        ws[ws.max_row][0].font = BOLD
        start = ws.max_row + 1
        ws.append([c for c, _ in ISSUE_COLUMNS])
        for i, cell in enumerate(ws[ws.max_row], start=1):
            cell.fill, cell.font, cell.border, cell.alignment = HEADER, BOLD, BORDER, WRAP
            ws.column_dimensions[get_column_letter(i)].width = max(ws.column_dimensions[get_column_letter(i)].width or 0, ISSUE_COLUMNS[i - 1][1])
        for issue in r["issues"]:
            ws.append(issue_values(r["round"], issue))
            for cell in ws[ws.max_row]:
                cell.alignment, cell.border = WRAP, BORDER
        if r.get("praise"):
            ws.append([])
            ws.append(["What the tester liked"])
            ws[ws.max_row][0].font = BOLD
            for p in r["praise"]:
                ws.append([p])
        ws.freeze_panes = ws.cell(row=start + 1, column=1)

    wb.save(TARGET)
    print(f"wrote {TARGET} ({len(rounds)} rounds)")


if __name__ == "__main__":
    build()
